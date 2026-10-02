"""Literal OOXML facts. Never evaluate a formula or invent document page numbers."""

from __future__ import annotations

import mimetypes
import posixpath
from collections import defaultdict
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docgrain_domain.source_format import SourceFormat

from .structural import StructuralItem, VerifiedSource, _issue

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
}
W = "{" + NS["w"] + "}"
R = "{" + NS["r"] + "}"


def relationships(package: ZipFile, part: str) -> dict:
    name = posixpath.join(
        posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels"
    )
    if name not in package.namelist():
        return {}
    result = {}
    for rel in ET.fromstring(package.read(name)):
        external = rel.get("TargetMode") == "External"
        target = rel.get("Target", "")
        resolved = posixpath.normpath(
            target.lstrip("/")
            if target.startswith("/")
            else posixpath.join(posixpath.dirname(part), target)
        )
        result[rel.get("Id")] = {
            "type": rel.get("Type", "").rsplit("/", 1)[-1],
            "target": target if external else resolved,
            "external": external,
        }
    return result


def _children(element, path):
    counts = defaultdict(int)
    for child in element:
        tag = child.tag.rsplit("}", 1)[-1]
        counts[tag] += 1
        yield child, f"{path}/{tag}[{counts[tag]}]"


def _text(element):
    # Text boxes have their own blocks. Deleted text/field instructions are not visible text.
    if element.tag in {W + "txbxContent", W + "del", W + "instrText"}:
        return ""
    if element.tag == W + "t":
        return element.text or ""
    if element.tag == W + "tab":
        return "\t"
    if element.tag in {W + "br", W + "cr"}:
        return "\n"
    return "".join(_text(child) for child in element)


def docx_items(
    source: VerifiedSource, issues: list
) -> tuple[list[StructuralItem], dict]:
    items = []
    metadata = {
        "parts": [],
        "relationships": [],
        "fields": [],
        "drawings": [],
        "paragraphs": [],
        "sections": [],
        "order": "body_then_referenced_parts; not rendered page order",
    }
    with ZipFile(source.path) as package:
        main = "word/document.xml"
        root = ET.fromstring(package.read(main))
        rels = relationships(package, main)
        for section in root.iter(W + "sectPr"):
            metadata["sections"].append(
                {
                    "references": [
                        {
                            "kind": e.tag.rsplit("}", 1)[-1],
                            "type": e.get(W + "type"),
                            "relationship_id": e.get(R + "id"),
                            "target": rels.get(e.get(R + "id"), {}).get("target"),
                        }
                        for e in section
                        if e.tag in {W + "headerReference", W + "footerReference"}
                    ]
                }
            )
        styles = {}
        if "word/styles.xml" in package.namelist():
            for style in ET.fromstring(package.read("word/styles.xml")).findall(
                "w:style", NS
            ):
                outline = style.find("w:pPr/w:outlineLvl", NS)
                if outline is not None:
                    styles[style.get(W + "styleId")] = (
                        int(outline.get(W + "val", "0")) + 1
                    )
        parts = [(main, root.find("w:body", NS), "/document/body")]
        seen = {main}
        # Only referenced parts; unused header copies are not source content.
        for element in root.iter():
            if element.tag in {
                W + "headerReference",
                W + "footerReference",
                W + "footnoteReference",
                W + "endnoteReference",
            }:
                if element.tag in {W + "footnoteReference", W + "endnoteReference"}:
                    rel = next(
                        (
                            v
                            for v in rels.values()
                            if v["type"]
                            == element.tag.rsplit("}", 1)[-1].replace("Reference", "s")
                        ),
                        None,
                    )
                else:
                    rel = rels.get(element.get(R + "id"))
                if (
                    not rel
                    or rel["external"]
                    or rel["target"] not in package.namelist()
                ):
                    issues.append(
                        _issue(
                            SourceFormat.DOCX,
                            "part_reference_unresolved",
                            "Referenced header/footer/note part is missing or external",
                            location=main,
                            item_ref=element.get(R + "id") or element.get(W + "id"),
                        )
                    )
                if rel and not rel["external"] and rel["target"] not in seen:
                    target = rel["target"]
                    if target in package.namelist():
                        part_root = ET.fromstring(package.read(target))
                        parts.append(
                            (target, part_root, "/" + part_root.tag.rsplit("}", 1)[-1])
                        )
                        seen.add(target)
        referenced_notes = {
            kind: {e.get(W + "id") for e in root.iter(W + kind + "Reference")}
            for kind in ("footnote", "endnote")
        }

        def locator(part, path):
            return {"kind": "docx_block", "part": part, "path": path}

        def drawings(element, part, path):
            for child, child_path in _children(element, path):
                if child.tag.rsplit("}", 1)[-1] == "docPr":
                    metadata["drawings"].append(
                        {
                            "part": part,
                            "path": child_path,
                            "title": child.get("title"),
                            "description": child.get("descr"),
                        }
                    )
                if child.tag == W + "txbxContent":
                    walk(child, part, child_path)
                elif child.tag == "{" + NS["a"] + "}blip":
                    rid = child.get(R + "embed") or child.get(R + "link")
                    rel = relationships(package, part).get(rid)
                    if (
                        not rel
                        or rel["external"]
                        or rel["target"] not in package.namelist()
                    ):
                        issues.append(
                            _issue(
                                SourceFormat.DOCX,
                                "drawing_binary_unavailable",
                                "Linked/missing image; binary not invented",
                                location=f"{part}:{child_path}",
                            )
                        )
                    else:
                        data = package.read(rel["target"])
                        items.append(
                            StructuralItem(
                                "picture",
                                f"{part}:{child_path}",
                                locator(part, child_path),
                                asset_bytes=data,
                                asset_mime=mimetypes.guess_type(rel["target"])[0]
                                or "application/octet-stream",
                            )
                        )
                        metadata["relationships"].append(
                            {
                                "part": part,
                                "path": child_path,
                                "relationship_id": rid,
                                "target": rel["target"],
                            }
                        )
                elif child.tag == "{" + NS["c"] + "}chart":
                    issues.append(
                        _issue(
                            SourceFormat.DOCX,
                            "docx_chart_unextracted",
                            "Embedded Word chart not extracted by this profile",
                            location=f"{part}:{child_path}",
                        )
                    )
                else:
                    drawings(child, part, child_path)

        def table(element, part, path):
            rows = []
            active = {}
            for row_index, (row, row_path) in enumerate(
                v for v in _children(element, path) if v[0].tag == W + "tr"
            ):
                cells = []
                before = row.find("w:trPr/w:gridBefore", NS)
                col = int(before.get(W + "val", "0")) if before is not None else 0
                cells.extend({"value": None} for _ in range(col))
                for cell, cell_path in (
                    v for v in _children(row, row_path) if v[0].tag == W + "tc"
                ):
                    span = cell.find("w:tcPr/w:gridSpan", NS)
                    width = int(span.get(W + "val", "1")) if span is not None else 1
                    merge = cell.find("w:tcPr/w:vMerge", NS)
                    continued = (
                        merge is not None
                        and merge.get(W + "val", "continue") == "continue"
                    )
                    value = "\n".join(_text(p) for p in cell.findall("w:p", NS))
                    attrs = {
                        "part": part,
                        "path": cell_path,
                        "merge_covered": continued,
                    }
                    mapped = {
                        "value": value,
                        "row_span": 1,
                        "col_span": width,
                        "locator": locator(part, cell_path),
                        "source_attributes": attrs,
                    }
                    if continued:
                        parent = active.get(col)
                        if parent is None:
                            issues.append(
                                _issue(
                                    SourceFormat.DOCX,
                                    "merge_without_origin",
                                    "Vertical merge has no preceding origin",
                                    location=f"{part}:{cell_path}",
                                )
                            )
                        else:
                            parent["row_span"] += 1
                        mapped["value"] = None
                    elif merge is not None:
                        active[col] = mapped
                    else:
                        active.pop(col, None)
                    cells.append(mapped)
                    cells.extend(
                        {
                            "value": None,
                            "source_attributes": {"merge_covered": True},
                            "locator": locator(part, cell_path),
                        }
                        for _ in range(width - 1)
                    )
                    col += width
                    drawings(cell, part, cell_path)
                    for nested, nested_path in _children(cell, cell_path):
                        if nested.tag == W + "tbl":
                            table(nested, part, nested_path)
                rows.append(cells)
            if rows:
                width = max(map(len, rows))
                for row in rows:
                    row.extend({"value": None} for _ in range(width - len(row)))
                items.append(
                    StructuralItem(
                        "table", f"{part}:{path}", locator(part, path), cells=rows
                    )
                )

        def walk(element, part, path):
            if element is None:
                return
            for child, child_path in _children(element, path):
                tag = child.tag.rsplit("}", 1)[-1]
                if tag == "p":
                    value = _text(child)
                    style = child.find("w:pPr/w:pStyle", NS)
                    outline = child.find("w:pPr/w:outlineLvl", NS)
                    level = (
                        int(outline.get(W + "val", "0")) + 1
                        if outline is not None
                        else styles.get(
                            style.get(W + "val") if style is not None else None
                        )
                    )
                    if level == 10:
                        level = None  # OOXML outlineLvl=9 explicitly means body text.
                    metadata["paragraphs"].append(
                        {
                            "part": part,
                            "path": child_path,
                            "style": style.get(W + "val")
                            if style is not None
                            else None,
                            "outline_level": level,
                            "numbering": {
                                e.tag.rsplit("}", 1)[-1]: e.get(W + "val")
                                for e in child.findall("w:pPr/w:numPr/*", NS)
                            },
                        }
                    )
                    if child.find(".//w:del", NS) is not None:
                        issues.append(
                            _issue(
                                SourceFormat.DOCX,
                                "tracked_changes_unreviewed",
                                "Deleted text is not current visible text; tracked changes require review",
                                location=f"{part}:{child_path}",
                            )
                        )
                    kind = (
                        "heading"
                        if level
                        else "list_item"
                        if child.find("w:pPr/w:numPr", NS) is not None
                        or (
                            style is not None
                            and style.get(W + "val", "").startswith("List")
                        )
                        else "paragraph"
                    )
                    if value:
                        items.append(
                            StructuralItem(
                                kind,
                                f"{part}:{child_path}",
                                locator(part, child_path),
                                text=value,
                                level=level or 1,
                            )
                        )
                    for field in child.findall(".//w:instrText", NS):
                        metadata["fields"].append(
                            {
                                "part": part,
                                "path": child_path,
                                "instruction": field.text or "",
                                "visible_result": value,
                            }
                        )
                        issues.append(
                            _issue(
                                SourceFormat.DOCX,
                                "field_result_unverified",
                                "Stored field result retained; field not evaluated",
                                location=f"{part}:{child_path}",
                            )
                        )
                    for field in child.findall(".//w:fldSimple", NS):
                        metadata["fields"].append(
                            {
                                "part": part,
                                "path": child_path,
                                "instruction": field.get(W + "instr", ""),
                                "visible_result": _text(field),
                            }
                        )
                        issues.append(
                            _issue(
                                SourceFormat.DOCX,
                                "field_result_unverified",
                                "Stored simple field result retained; not evaluated",
                                location=f"{part}:{child_path}",
                            )
                        )
                    drawings(child, part, child_path)
                elif tag == "tbl":
                    table(child, part, child_path)
                elif tag in {"footnote", "endnote"}:
                    if child.get(W + "id") in referenced_notes[tag]:
                        walk(child, part, child_path)
                elif tag in {"sdt", "sdtContent", "ins"}:
                    walk(child, part, child_path)
                elif tag in {"altChunk", "object", "del"}:
                    issues.append(
                        _issue(
                            SourceFormat.DOCX,
                            "unsupported_docx_content",
                            f"Unresolved {tag} content",
                            location=f"{part}:{child_path}",
                        )
                    )
            if any(e.tag == W + "ins" for e in element.iter()):
                issues.append(
                    _issue(
                        SourceFormat.DOCX,
                        "tracked_changes_unreviewed",
                        "Inserted text retained; tracked changes require review",
                        location=part,
                    )
                )

        for part, element, path in parts:
            metadata["parts"].append(part)
            walk(element, part, path)
    return items, metadata


def xlsx_charts(
    source: VerifiedSource, workbook, cached, issues: list
) -> list[StructuralItem]:
    """Read chart XML and resolve local series references against native cells, retaining caches."""
    import re

    from openpyxl.utils.cell import range_boundaries

    items = []
    with ZipFile(source.path) as package:
        workbook_part = "xl/workbook.xml"
        rels = relationships(package, workbook_part)
        for sheet in ET.fromstring(package.read(workbook_part)).findall(
            "s:sheets/s:sheet", NS
        ):
            name = sheet.get("name")
            sheet_part = rels[sheet.get(R + "id")]["target"]
            for drawing in ET.fromstring(package.read(sheet_part)).findall(
                "s:drawing", NS
            ):
                drawing_rel = relationships(package, sheet_part).get(
                    drawing.get(R + "id")
                )
                if not drawing_rel or drawing_rel["external"]:
                    continue
                drawing_part = drawing_rel["target"]
                drawing_rels = relationships(package, drawing_part)
                for index, anchor in enumerate(
                    ET.fromstring(package.read(drawing_part))
                ):
                    chart = anchor.find(".//c:chart", NS)
                    if chart is None:
                        continue
                    rel = drawing_rels.get(chart.get(R + "id"))
                    if (
                        not rel
                        or rel["external"]
                        or rel["target"] not in package.namelist()
                    ):
                        issues.append(
                            _issue(
                                SourceFormat.XLSX,
                                "chart_part_unavailable",
                                "Chart part is external/missing",
                                location=name,
                            )
                        )
                        continue
                    part = rel["target"]
                    xml = ET.fromstring(package.read(part))
                    origin = anchor.find("xdr:from", NS)
                    if origin is None:
                        issues.append(
                            _issue(
                                SourceFormat.XLSX,
                                "chart_anchor_missing",
                                "Absolute chart anchor has no source cell",
                                location=part,
                            )
                        )
                        continue
                    col = int(origin.findtext("xdr:col", "0", NS)) + 1
                    row = int(origin.findtext("xdr:row", "0", NS)) + 1
                    from openpyxl.utils import get_column_letter

                    a1 = f"{get_column_letter(col)}{row}"
                    locator = {
                        "kind": "spreadsheet_range",
                        "sheet": name,
                        "a1_range": a1,
                    }
                    data = {
                        "format": "docgrain.native-chart",
                        "version": "1.0.0",
                        "part": part,
                        "drawing_part": drawing_part,
                        "title": "".join(
                            t.text or ""
                            for t in xml.findall("c:chart/c:title//a:t", NS)
                        ),
                        "types": [],
                        "series": [],
                        "axes": [],
                    }
                    bindings = {}

                    def vector(element, key, bindings=bindings, part=part):
                        formula = element.findtext(".//c:f", None, NS)
                        points = [
                            {
                                "index": int(p.get("idx", "0")),
                                "value": p.findtext("c:v", "", NS),
                            }
                            for p in element.findall(".//c:pt", NS)
                        ]
                        result = {
                            "formula": formula,
                            "stored_points": points,
                            "cells": [],
                            "locator": None,
                            "format_code": element.findtext(
                                ".//c:formatCode", None, NS
                            ),
                        }
                        if formula:
                            match = re.fullmatch(
                                r"(?:'((?:[^']|'')+)'|([^'!]+))!([$A-Za-z0-9:]+)",
                                formula,
                            )
                            try:
                                if not match:
                                    raise ValueError("unsupported or external range")
                                sheet_name = (match[1] or match[2]).replace("''", "'")
                                if sheet_name not in workbook.sheetnames:
                                    raise ValueError("missing source sheet")
                                bounds = range_boundaries(match[3].replace("$", ""))
                                x0, y0, x1, y1 = bounds
                                if (
                                    None in bounds
                                    or x0 < 1
                                    or y0 < 1
                                    or x1 < x0
                                    or y1 < y0
                                    or x1 > 16384
                                    or y1 > 1048576
                                    or (x1 - x0 + 1) * (y1 - y0 + 1) > 10000
                                ):
                                    raise ValueError("range exceeds 10000 cells")
                                source_locator = {
                                    "kind": "spreadsheet_range",
                                    "sheet": sheet_name,
                                    "a1_range": match[3].replace("$", "").upper(),
                                }
                                result["locator"] = source_locator
                                bindings[key] = source_locator
                                for cells in workbook[sheet_name].iter_rows(
                                    min_row=y0, max_row=y1, min_col=x0, max_col=x1
                                ):
                                    for cell in cells:
                                        formula_cell = (
                                            cell.value
                                            if cell.data_type == "f"
                                            else None
                                        )
                                        value = (
                                            cached[sheet_name][cell.coordinate].value
                                            if formula_cell
                                            else cell.value
                                        )
                                        if hasattr(value, "isoformat"):
                                            value = value.isoformat()
                                        result["cells"].append(
                                            {
                                                "coordinate": cell.coordinate,
                                                "value": value,
                                                "formula": formula_cell,
                                                "number_format": cell.number_format,
                                                "data_type": cell.data_type,
                                            }
                                        )
                                        if formula_cell and value is None:
                                            issues.append(
                                                _issue(
                                                    SourceFormat.XLSX,
                                                    "chart_formula_cache_missing",
                                                    "Chart source formula has no cached result; not evaluated",
                                                    location=f"{sheet_name}!{cell.coordinate}",
                                                )
                                            )
                            except (ValueError, TypeError):
                                issues.append(
                                    _issue(
                                        SourceFormat.XLSX,
                                        "chart_reference_unresolved",
                                        "Chart reference is external, unsupported or missing; stored cache is not current source authority",
                                        location=part,
                                        item_ref=formula,
                                    )
                                )
                        return result

                    plot = xml.find("c:chart/c:plotArea", NS)
                    for child in plot if plot is not None else []:
                        tag = child.tag.rsplit("}", 1)[-1]
                        if tag.endswith("Chart"):
                            data["types"].append(tag)
                            for ser in child.findall("c:ser", NS):
                                series = {
                                    "index": int(ser.find("c:idx", NS).get("val", "0")),
                                    "type": tag,
                                }
                                for child_name in (
                                    "tx",
                                    "cat",
                                    "val",
                                    "xVal",
                                    "yVal",
                                    "bubbleSize",
                                ):
                                    entry = ser.find("c:" + child_name, NS)
                                    if entry is not None:
                                        key = f"/source_data/series/{len(data['series'])}/{child_name}"
                                        series[child_name] = vector(entry, key)
                                        if child_name == "tx":
                                            series[child_name]["literal"] = (
                                                entry.findtext("c:v", None, NS)
                                            )
                                data["series"].append(series)
                        elif tag.endswith("Ax"):
                            data["axes"].append(
                                {
                                    "type": tag,
                                    "title": "".join(
                                        t.text or ""
                                        for t in child.findall("c:title//a:t", NS)
                                    ),
                                    "number_format": child.find("c:numFmt", NS).get(
                                        "formatCode"
                                    )
                                    if child.find("c:numFmt", NS) is not None
                                    else None,
                                }
                            )
                    if not data["series"]:
                        issues.append(
                            _issue(
                                SourceFormat.XLSX,
                                "chart_data_unavailable",
                                "Unsupported chart or no native series",
                                location=part,
                            )
                        )
                    if xml.find("c:externalData", NS) is not None:
                        issues.append(
                            _issue(
                                SourceFormat.XLSX,
                                "chart_external_data",
                                "Chart uses external data; current values not verified",
                                location=part,
                            )
                        )
                    issues.append(
                        _issue(
                            SourceFormat.XLSX,
                            "chart_visual_unverified",
                            "Native series preserved; rendering, trend/visual interpretation not verified",
                            location=part,
                        )
                    )
                    items.append(
                        StructuralItem(
                            "chart",
                            f"{drawing_part}:anchor:{index}:{part}",
                            locator,
                            source_data=data,
                            field_locators=bindings,
                        )
                    )
    return items
