"""Offline source comparisons. None of these checks certify complete document meaning."""

import re
import unicodedata
from collections import Counter
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

from docgrain_domain.canonical import CanonicalKnowledgeSnapshot
from docgrain_domain.canonical.identity import canonical_json_bytes


def normalized_text(value):
    return " ".join(unicodedata.normalize("NFKC",str(value)).split())


def node_text(node):
    if node.kind == "table":
        return " ".join(str(cell.value) for row in node.rows for cell in row if cell.value is not None)
    return getattr(node,"text",None) or getattr(node,"heading",None) or ""


def golden_table_checks(snapshot, checks):
    """Externally reviewed cell/column labels; never derive expectations from canonical."""
    nodes = {node.id:node for node in snapshot.structure}
    results = []
    for check in checks:
        if check["revision_id"] != snapshot.knowledge_revision.id or check["source_sha256"] != snapshot.source_version.content_sha256:
            raise ValueError("golden assertion belongs to another source/revision")
        row,column = check["row"],check["column"]
        if row < 0 or column < 0:
            raise ValueError("golden cell coordinates must be nonnegative")
        node = nodes.get(check["node_id"])
        actual = None
        found = bool(node and node.kind == "table" and row < len(node.rows) and column < len(node.rows[row]))
        if found:
            actual = node.rows[row][column].value
        results.append({**check,"found":found,"actual":actual,
                        "matches":found and actual == check["expected"]})
    return results


def audit_source(snapshot: CanonicalKnowledgeSnapshot, path: Path, *, golden_checks=()):
    snapshot = CanonicalKnowledgeSnapshot.model_validate(snapshot.model_dump(mode="json"))
    data = path.read_bytes()
    source = snapshot.source_version
    if len(data) != source.byte_size or sha256(data).hexdigest() != source.content_sha256:
        raise ValueError("audit source checksum/size differs from canonical source")
    nodes = snapshot.structure
    output = {"format":"docgrain.source-fidelity-review","version":"1.0.0",
              "document_id":snapshot.document_id,"revision_id":snapshot.knowledge_revision.id,
              "source_sha256":source.content_sha256,"source_bytes_verified":True,
              "canonical_sha256":sha256(canonical_json_bytes(snapshot.model_dump(mode="json"))).hexdigest(),
              "semantic_acceptance":"not_certified", "scope":[],"checks":{}}
    extension = Path(source.filename).suffix.lower()
    if extension == ".txt":
        text = data.decode("utf-8-sig")
        actual = "\n".join(node_text(node) for node in nodes if node_text(node))
        output["scope"] = ["UTF-8 text after NFKC/whitespace normalization; no semantic or ordering inference"]
        output["checks"] = {"source_characters":len(text),"canonical_characters":len(actual),
                            "normalized_equal":normalized_text(text)==normalized_text(actual)}
    elif extension == ".docx":
        with ZipFile(BytesIO(data)) as archive:
            document = ElementTree.fromstring(archive.read("word/document.xml"))
        ns = {"w":"http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        paragraphs = [normalized_text("".join(t.text or "" for t in p.findall(".//w:t",ns)))
                      for p in document.findall(".//w:body//w:p",ns)]
        expected = Counter(p for p in paragraphs if p)
        actual = Counter(normalized_text(node_text(n)) for n in nodes if node_text(n))
        output["scope"] = ["Main document OOXML paragraph text; headers/footers/drawing meaning/styles not assessed"]
        output["checks"] = {"source_paragraphs":sum(expected.values()),
                            "matched_occurrences":sum((expected&actual).values()),"missing":dict(expected-actual)}
    elif extension == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(data),data_only=False)
        cached = load_workbook(BytesIO(data),data_only=True)
        evidence = {e.id:e.locator for e in snapshot.evidence}
        actual = {}
        for node in nodes:
            for row in getattr(node,"rows",[]):
                for cell in row:
                    for eid in cell.annotation.provenance.evidence_ids if cell.annotation else []:
                        loc = evidence[eid]
                        if loc.kind == "spreadsheet_range":
                            actual.setdefault((loc.sheet,loc.a1_range),[]).append(cell)
        expected_count = 0
        errors = []
        def json_value(value):
            return value.isoformat() if hasattr(value,"isoformat") else value
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    if cell.value is None:
                        continue
                    expected_count += 1
                    formula = cell.value if cell.data_type == "f" else None
                    value = json_value(cached[sheet.title][cell.coordinate].value if formula else cell.value)
                    merged = next((area for area in sheet.merged_cells.ranges if cell.coordinate in area),None)
                    facts = {"value":value,"formula":formula,
                             "cached_value":json_value(cached[sheet.title][cell.coordinate].value) if formula else None,
                             "row_span":merged.max_row-merged.min_row+1 if merged and cell.row==merged.min_row and cell.column==merged.min_col else 1,
                             "col_span":merged.max_col-merged.min_col+1 if merged and cell.row==merged.min_row and cell.column==merged.min_col else 1}
                    candidates = actual.get((sheet.title,cell.coordinate),[])
                    if len(candidates)!=1 or any(getattr(candidates[0],k)!=v for k,v in facts.items()):
                        errors.append({"sheet":sheet.title,"cell":cell.coordinate,"expected":facts,
                                       "actual":[{k:getattr(c,k) for k in facts} for c in candidates]})
        workbook.close()
        cached.close()
        output["scope"] = ["Nonempty spreadsheet values/formulas/cache/merged spans at exact source coordinates; visual/style meaning not assessed"]
        output["checks"] = {"source_nonempty_cells":expected_count,"mapped_coordinates":len(actual),"mismatches":errors}
    elif extension == ".pdf":
        import pymupdf
        evidence = {e.id:e.locator for e in snapshot.evidence}
        def tokens(value):
            return Counter(re.findall(r"\w+",unicodedata.normalize("NFKC",value).casefold()))
        pages = []
        with pymupdf.open(stream=data,filetype="pdf") as document:
            for page in document:
                number = page.number+1
                expected = tokens(page.get_text())
                actual = tokens("\n".join(node_text(n) for n in nodes if n.kind!="document" and any(
                    getattr(evidence[e],"page_number",None)==number for e in n.annotation.provenance.evidence_ids)))
                pages.append({"page_number":number,"native_token_occurrences":sum(expected.values()),
                              "matched_occurrences":sum((expected&actual).values()),"unmatched":dict(expected-actual)})
        output["scope"] = ["PDF native token occurrence overlap only; segmentation, cell semantics, reading order and visual meaning require review"]
        output["checks"] = {"pages":pages}
    else:
        raise ValueError("source audit format unsupported")
    output["golden_table_checks"] = golden_table_checks(snapshot,golden_checks)
    output["limitations"] = ["A valid JSON or high token overlap does not establish semantic completeness.",
                              "This audit does not clear visual gaps or mutate any canonical snapshot."]
    return output
