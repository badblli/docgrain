"""Small neutral corpus for the profile benchmark; never company source data."""

from datetime import date
from pathlib import Path


def create_profile_corpus(root: Path) -> Path:
    import pymupdf
    from docx import Document
    from openpyxl import Workbook

    company = root / "synthetic"
    company.mkdir(parents=True, exist_ok=True)
    pdf = pymupdf.open()
    page = pdf.new_page(width=600, height=800)
    for y in range(80, 150, 15):
        page.insert_text((60, y), f"Left column {y}")
        page.insert_text((330, y), f"Right column {y}")
    page = pdf.new_page(width=600, height=800)
    for x in (60, 200, 340):
        page.draw_line((x, 100), (x, 190))
    for y in (100, 145, 190):
        page.draw_line((60, y), (340, y))
    for x, y, text in ((70, 125, "Item"), (210, 125, "Count"),
                       (70, 170, "Sample"), (210, 170, "42")):
        page.insert_text((x, y), text)
    pdf.save(company / "columns-table.pdf")
    pdf.close()
    printed = Path(__file__).with_name("printed-tr-en.png")
    (company / "scan.png").write_bytes(printed.read_bytes())
    pdf = pymupdf.open()
    page = pdf.new_page(width=750, height=250)
    page.insert_image(page.rect, filename=str(printed))
    pdf.save(company / "scan.pdf")
    pdf.close()
    doc = Document()
    doc.add_heading("Synthetic document", 1)
    paragraph = doc.add_paragraph()
    paragraph.add_run("Combined")
    paragraph.add_run("Runs")
    doc.add_paragraph("Example text 42")
    doc.save(company / "sample.docx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Percent", "Date"])
    sheet.append([0.15, date(2026, 1, 2)])
    sheet["A2"].number_format = "0%"
    sheet["B2"].number_format = "yyyy-mm-dd"
    workbook.save(company / "sample.xlsx")
    workbook.close()
    (company / "sample.txt").write_text("# Heading\nExample text 42\n", encoding="utf-8")
    return root
