"""Generate redistributable M1b documents in a temporary directory."""

from __future__ import annotations

from pathlib import Path

import pymupdf
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image


def create_corpus(root: Path) -> dict[str, Path]:
    root.mkdir(parents=True, exist_ok=True)
    files: dict[str, Path] = {}
    printed = Path(__file__).with_name("printed-tr-en.png")
    files["png-printed"] = printed
    with Image.open(printed) as original_image:
        exif = Image.Exif()
        exif[274] = 6
        jpeg = root / "printed-exif.jpg"
        original_image.transpose(Image.Transpose.ROTATE_90).save(jpeg, quality=98, exif=exif)
        files["jpeg-printed"] = jpeg
    image = root / "red-square.png"
    Image.new("RGB", (40, 40), (220, 40, 40)).save(image)

    def pdf(name: str, *, rotation: int = 0, crop: bool = False, text: bool = True,
            columns: bool = False, picture: bool = False, table: bool = False) -> None:
        path = root / f"{name}.pdf"
        document = pymupdf.open()
        page = document.new_page(width=600, height=800)
        if text:
            page.insert_text((72, 72), "Synthetic heading", fontsize=18)
            page.insert_text((72, 110), "Deterministic paragraph.")
        if columns:
            page.insert_text((320, 110), "Second column.")
        if table:
            for x in (72, 180, 300):
                page.draw_line((x, 150), (x, 240))
            for y in (150, 195, 240):
                page.draw_line((72, y), (300, y))
            page.insert_text((85, 175), "Item")
            page.insert_text((190, 175), "Count")
            page.insert_text((85, 220), "A")
            page.insert_text((190, 220), "2")
        if picture:
            page.insert_image(pymupdf.Rect(72, 280, 192, 400), filename=str(image))
        if crop:
            page.set_cropbox(pymupdf.Rect(20, 20, 580, 760))
        page.set_rotation(rotation)
        document.set_metadata({"title": name, "creator": "Docgrain synthetic corpus"})
        document.save(path, no_new_id=True)
        document.close()
        files[name] = path

    pdf("basic")
    pdf("table", table=True)
    pdf("multicolumn", columns=True)
    pdf("rotated90", rotation=90)
    pdf("rotated180", rotation=180)
    pdf("rotated270", rotation=270)
    pdf("cropped", crop=True)
    pdf("image-heavy", text=False, picture=True)
    pdf("scanned-low-text", text=False)

    doc = Document()
    doc.add_heading("Heading One", 1)
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("List entry", style="List Bullet")
    doc.add_picture(str(image), width=Inches(0.5))
    path = root / "headings-list-image.docx"
    doc.save(path)
    files["docx-headings"] = path
    doc = Document()
    doc.add_heading("Table section", 1)
    table_doc = doc.add_table(rows=2, cols=2)
    for row, values in zip(table_doc.rows, (("Name", "Count"), ("A", "2"))):
        for cell, value in zip(row.cells, values):
            cell.text = value
    path = root / "table.docx"
    doc.save(path)
    files["docx-table"] = path

    for name, value in {
        "utf8": "# Heading\nFirst line\nSecond line\n",
        "bom": "\ufeff# Heading\r\nText\r\n",
        "multilingual": "# Başlık\nİzmir 🌊\n日本語\n",
    }.items():
        path = root / f"{name}.txt"
        path.write_bytes(value.encode("utf-8"))
        files[f"txt-{name}"] = path

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Count", "Double"])
    sheet.append(["A", 2, "=B2*2"])
    sheet.merge_cells("A4:B4")
    sheet["A4"] = "Merged"
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    sheet.add_chart(chart, "E2")
    sheet.add_image(SheetImage(str(image)), "G2")
    workbook.create_sheet("Second")["A1"] = "Other sheet"
    path = root / "workbook.xlsx"
    workbook.save(path)
    files["xlsx"] = path
    return files
