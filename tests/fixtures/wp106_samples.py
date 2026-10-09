"""Small synthetic WP106 samples, generated in memory; no customer files, no stored binaries."""

from __future__ import annotations

import io
import json
import wave
import zipfile
from email.message import EmailMessage

from PIL import Image, ImageDraw

TEXT = "Garden room sleeps two guests"


def _zip(members: list[tuple[str, str | bytes]], *, stored_first: bool = False) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for index, (name, content) in enumerate(members):
            kind = zipfile.ZIP_STORED if stored_first and index == 0 else zipfile.ZIP_DEFLATED
            archive.writestr(zipfile.ZipInfo(name), content, compress_type=kind)
    return buffer.getvalue()


def pptx() -> bytes:
    try:
        from pptx import Presentation
    except ImportError:  # host without python-pptx: a minimal package is enough for sniffing
        return _zip([("[Content_Types].xml", "<Types/>"), ("ppt/presentation.xml", "<p:presentation/>")])
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Rooms"
    slide.placeholders[1].text = TEXT
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def html() -> bytes:
    return (f"<!DOCTYPE html><html><head><title>Rooms</title></head><body><h1>Rooms</h1><p>{TEXT}</p>"
            "<table><tr><th>Room</th><th>Guests</th></tr><tr><td>Garden</td><td>2</td></tr></table>"
            "</body></html>").encode()


def markdown() -> bytes:
    return f"# Rooms\n\n{TEXT}.\n\n| Room | Guests |\n|---|---|\n| Garden | 2 |\n".encode()


def csv() -> bytes:
    return b"Room,Guests\nGarden,2\nSea view,3\n"


def odt() -> bytes:
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">'
        f'<office:body><office:text><text:h text:outline-level="1">Rooms</text:h><text:p>{TEXT}</text:p>'
        '</office:text></office:body></office:document-content>')
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2">'
        '<manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/>'
        '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
        '</manifest:manifest>')
    return _zip([("mimetype", "application/vnd.oasis.opendocument.text"), ("content.xml", content),
                 ("META-INF/manifest.xml", manifest)], stored_first=True)


def eml() -> bytes:
    message = EmailMessage()
    message["From"] = "front.desk@example.test"
    message["To"] = "guest@example.test"
    message["Subject"] = "Room details"
    message["Date"] = "Fri, 09 Oct 2026 10:00:00 +0000"
    message.set_content(f"{TEXT}.\nBreakfast starts at 07:00.")
    return bytes(message)


def epub() -> bytes:
    container = ('<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                 '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
                 '</rootfiles></container>')
    opf = ('<?xml version="1.0" encoding="UTF-8"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" '
           'unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
           '<dc:identifier id="id">urn:example:rooms</dc:identifier><dc:title>Rooms</dc:title><dc:language>en</dc:language>'
           '</metadata><manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>'
           '<item id="c1" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest>'
           '<spine><itemref idref="c1"/></spine></package>')
    nav = ('<?xml version="1.0" encoding="UTF-8"?><html xmlns="http://www.w3.org/1999/xhtml" '
           'xmlns:epub="http://www.idpf.org/2007/ops"><head><title>Nav</title></head><body>'
           '<nav epub:type="toc"><ol><li><a href="chapter.xhtml">Rooms</a></li></ol></nav></body></html>')
    chapter = ('<?xml version="1.0" encoding="UTF-8"?><html xmlns="http://www.w3.org/1999/xhtml"><head>'
               f'<title>Rooms</title></head><body><h1>Rooms</h1><p>{TEXT}</p></body></html>')
    return _zip([("mimetype", "application/epub+zip"), ("META-INF/container.xml", container),
                 ("OEBPS/content.opf", opf), ("OEBPS/nav.xhtml", nav), ("OEBPS/chapter.xhtml", chapter)],
                stored_first=True)


def _page(text: str, size=(900, 300)) -> Image.Image:
    image = Image.new("RGB", size, "white")
    ImageDraw.Draw(image).text((40, 120), text, fill="black", font_size=48)
    return image


def tiff(pages: int = 2) -> bytes:
    images = [_page(f"Page {n} {TEXT}") for n in range(1, pages + 1)]
    buffer = io.BytesIO()
    images[0].save(buffer, format="TIFF", save_all=True, append_images=images[1:], compression="tiff_lzw")
    return buffer.getvalue()


def image(format_name: str) -> bytes:
    buffer = io.BytesIO()
    _page(TEXT).save(buffer, format=format_name, **({"lossless": True} if format_name == "WEBP" else {}))
    return buffer.getvalue()


def vtt() -> bytes:
    return f"WEBVTT\n\n00:00:00.000 --> 00:00:03.000\n{TEXT}.\n\n00:00:03.000 --> 00:00:06.000\nBreakfast at 07:00.\n".encode()


def jats() -> bytes:
    return ('<?xml version="1.0" encoding="UTF-8"?>\n'
            '<!DOCTYPE article PUBLIC "-//NLM//DTD JATS (Z39.96) Journal Publishing DTD v1.2 20190208//EN" '
            '"JATS-journalpublishing1.dtd">\n'
            '<article><front><article-meta><title-group><article-title>Rooms</article-title></title-group>'
            f'</article-meta></front><body><sec><title>Rooms</title><p>{TEXT}</p></sec></body></article>').encode()


def xbrl() -> bytes:
    return (b'<?xml version="1.0"?><xbrl xmlns="http://www.xbrl.org/2003/instance">'
            b'<context id="c"/></xbrl>')


def doclang() -> bytes:
    return b'<?xml version="1.0"?><doclang><text>Garden room</text></doclang>'


def docling_json() -> bytes:
    return json.dumps({"schema_name": "DoclingDocument", "version": "1.0.0", "name": "rooms"}).encode()


def wav() -> bytes:
    buffer = io.BytesIO()
    with wave.open(buffer, "wb") as audio:
        audio.setnchannels(1)
        audio.setsampwidth(2)
        audio.setframerate(8000)
        audio.writeframes(b"\x00\x00" * 800)
    return buffer.getvalue()


def legacy_doc() -> bytes:
    """Only the OLE header and a WordDocument stream name: enough for sniffing, not for reading."""
    return b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504 + "WordDocument".encode("utf-16-le") + b"\x00" * 64


SAMPLES = {
    "rooms.pptx": pptx, "rooms.html": html, "rooms.md": markdown, "rooms.csv": csv, "rooms.odt": odt,
    "rooms.eml": eml, "rooms.epub": epub, "rooms.tiff": tiff, "rooms.webp": lambda: image("WEBP"),
    "rooms.bmp": lambda: image("BMP"), "rooms.vtt": vtt, "rooms.xml": jats,
}
