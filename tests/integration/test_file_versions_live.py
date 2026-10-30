"""WP124 real PostgreSQL promotion and real offline DOCX conversion (worker image)."""

import importlib.util
import os
from contextlib import contextmanager
from uuid import uuid4

import psycopg
import pytest
from docgrain_api import document_jobs, repository
from docgrain_api.main import app
from docgrain_api.routers import documents
from docgrain_api.settings import get_settings
from docgrain_domain.canonical.version_diff import diff_docling_documents
from fastapi.testclient import TestClient
from psycopg import sql


@pytest.fixture
def database(monkeypatch):
    url = os.getenv("DOCGRAIN_M1_TEST_DATABASE_URL")
    if not url:
        pytest.skip("isolated PostgreSQL requires DOCGRAIN_M1_TEST_DATABASE_URL")
    schema = "wp124_test_" + uuid4().hex
    with psycopg.connect(url) as conn:
        conn.execute(sql.SQL("CREATE SCHEMA {}").format(sql.Identifier(schema)))

    @contextmanager
    def connection():
        with psycopg.connect(url, row_factory=psycopg.rows.dict_row) as conn:
            conn.execute(sql.SQL("SET search_path TO {}").format(sql.Identifier(schema)))
            yield conn

    monkeypatch.setattr(repository, "_connection", connection)
    monkeypatch.setattr(get_settings(), "use_fixtures", False)
    monkeypatch.setattr(documents.format_capabilities, "ensure_format_enabled", lambda fmt: fmt)
    monkeypatch.setattr(documents, "object_exists", lambda _: True)
    try:
        repository.initialize()
        repository.initialize()  # Additive columns also work on an existing database.
        yield connection
    finally:
        with psycopg.connect(url) as conn:
            conn.execute(sql.SQL("DROP SCHEMA {} CASCADE").format(sql.Identifier(schema)))


@pytest.mark.parametrize("completed_status", ["done", "partial"])
def test_api_version_promotes_only_after_usable_completion_and_never_rolls_back(database, completed_status):
    client = TestClient(app)
    payload = {"workspace_id": "ws_synthetic", "filename": "guide.txt", "mime_type": "text/plain",
               "byte_size": 10, "content_sha256": "a" * 64}
    original = client.post("/v1/documents", json=payload).json()
    document_id, v1 = original["document"]["id"], original["version"]["id"]
    with database() as conn:
        conn.execute("UPDATE document_versions SET status='done' WHERE id=%s", (v1,))
        conn.execute("UPDATE jobs SET status='done' WHERE id=%s", (original["job_id"],))
    path = f"/v1/documents/{document_id}/versions"
    second = client.post(path, json={**payload, "filename": "renamed.txt", "content_sha256": "b" * 64}).json()
    v2 = second["version"]["id"]
    assert second["version"]["revision"] == 2
    assert second["document"]["latest_version_id"] == v1
    with database() as conn:
        document_jobs.promote_usable_version(conn, v2)  # queued/processing does not promote
    assert client.get(f"/v1/documents/{document_id}").json()["latest_version"]["id"] == v1
    with database() as conn:
        conn.execute("UPDATE document_versions SET status=%s WHERE id=%s", (completed_status, v2))
        conn.execute("UPDATE jobs SET status=%s WHERE id=%s", (completed_status, second["job_id"]))
        document_jobs.promote_usable_version(conn, v2)
    assert client.get(f"/v1/documents/{document_id}").json()["latest_version"]["id"] == v2
    assert client.post(path, json={**payload, "content_sha256": "b" * 64}).json()["deduplicated"]
    third = client.post(path, json={**payload, "content_sha256": "c" * 64}).json()
    assert third["version"]["revision"] == 3
    with database() as conn:
        conn.execute("UPDATE document_versions SET status='failed' WHERE id=%s", (third["version"]["id"],))
        conn.execute("UPDATE jobs SET status='failed' WHERE id=%s", (third["job_id"],))
        document_jobs.promote_usable_version(conn, third["version"]["id"])
        document_jobs.promote_usable_version(conn, v1)  # late old completion cannot roll back
    assert client.get(f"/v1/documents/{document_id}").json()["latest_version"]["id"] == v2
    assert len(client.get(path).json()) == 3
    assert client.get(f"/v1/versions/{v1}").status_code == 200


@pytest.mark.skipif(importlib.util.find_spec("docling") is None, reason="Docling worker image required")
def test_real_docx_single_word_and_single_cell_diff(tmp_path, monkeypatch):
    from docling.datamodel.base_models import InputFormat
    from docling.document_converter import DocumentConverter
    from docx import Document

    monkeypatch.setenv("HF_HUB_OFFLINE", "1")
    monkeypatch.setenv("TRANSFORMERS_OFFLINE", "1")
    import requests
    monkeypatch.setattr(requests.sessions.Session, "send", lambda *a, **kw: pytest.fail("network forbidden"))
    paths = []
    for name, duration, price in (("v1", "three", "10"), ("v2", "four", "10"), ("v3", "four", "12")):
        document = Document()
        document.add_heading("Service prices", level=1)
        document.add_paragraph(f"Delivery takes {duration} days.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Service", "Price"
        table.cell(1, 0).text, table.cell(1, 1).text = "Delivery", price
        path = tmp_path / f"{name}.docx"
        document.save(path)
        paths.append(path)
    converter = DocumentConverter(allowed_formats=[InputFormat.DOCX])
    documents = [converter.convert(path).document for path in paths]
    change, = diff_docling_documents(documents[0], documents[1]).changes
    assert change.operation == "changed" and change.kind == "paragraph"
    assert change.before == "Delivery takes three days." and change.after == "Delivery takes four days."
    change, = diff_docling_documents(documents[1], documents[2]).changes
    assert change.operation == "changed" and change.kind == "table_cell"
    assert change.address["row_header"] == "Delivery" and change.address["column_header"] == "Price"
