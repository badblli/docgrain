"""Independent native facts, repeated Word blocks and geometry regression contracts."""

import json
from datetime import UTC, datetime
from hashlib import sha256
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import jsonschema
import pymupdf
import pytest
from docgrain_domain.canonical import SourceVersion
from docgrain_domain.canonical.ai_output import output_bundle, output_schema
from docgrain_domain.canonical.lifecycle import (
    processing_revision_id,
    source_revision_id,
)
from docgrain_domain.canonical.schema import generated_core_schema_text
from docgrain_domain.source_format import MIME_TYPES, SourceFormat
from docgrain_worker.canonical_mapper import CanonicalMapper
from docgrain_worker.canonical_writer import processing_spec
from docgrain_worker.native_office import docx_items, xlsx_charts
from docgrain_worker.pdf_fidelity import reconcile_pdf_tables
from docgrain_worker.structural import (
    StructuralItem,
    StructuralParseResult,
    VerifiedSource,
    _xlsx_cell,
)
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from PIL import Image


def verified(path):
    data = path.read_bytes()
    return VerifiedSource(path, sha256(data).hexdigest(), len(data))


def map_result(result, source):
    spec = processing_spec(result)
    now = datetime(2026, 10, 2, tzinfo=UTC)
    version = SourceVersion(
        id=source_revision_id("ws", "doc", source.content_sha256),
        document_id="doc",
        workspace_id="ws",
        content_sha256=source.content_sha256,
        byte_size=source.byte_size,
        filename=source.path.name,
        mime_type=MIME_TYPES[result.source_format],
        storage_uri="fixture://source",
        storage_version="v1",
        recorded_at=now,
    )
    for item in result.items:
        if item.asset_bytes:
            item.asset_path = "fixture://binary"
    return CanonicalMapper().map(
        result,
        version,
        revision_id=processing_revision_id(version.id, spec),
        created_at=now,
        processing=spec,
        pdf_path=source.path if result.source_format is SourceFormat.PDF else None,
    )


def result_for(fmt, items, issues):
    return StructuralParseResult(
        fmt,
        "docling",
        "2.130.0",
        "partial" if issues else "complete",
        items,
        ["source"],
        ["source"],
        issues,
        processing_options={
            "adapter_version": "n2-1",
            "canonical_schema_version": "0.6.0",
        },
    )


def office_file(root):
    image = root / "square.png"
    Image.new("RGB", (30, 30), "red").save(image)
    doc = Document()
    doc.add_heading("Exact section", 1)
    doc.add_paragraph("Repeat")
    p = doc.add_paragraph("Repeat")
    p.add_run().add_picture(str(image), width=Inches(0.2))
    # Same paragraph has text AND drawing; neither consumes the other's path.
    inline = p._p.xpath(".//wp:inline")[0]
    inline.tag = inline.tag.replace("inline", "anchor")
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Merged title"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 0).text = "Vertical"
    table.cell(1, 0).merge(table.cell(2, 0))
    table.cell(1, 1).text = "127"
    doc.sections[0].header.paragraphs[0].text = "Header exact 987"
    doc.sections[0].footer.paragraphs[0].text = "Footer exact 654"
    doc.sections[0].header.paragraphs[0].add_run().add_picture(
        str(image), width=Inches(0.2)
    )
    path = root / "parts.docx"
    doc.save(path)
    return path


def test_docx_real_parts_repeated_paragraphs_floating_drawing_and_merges(tmp_path):
    source = verified(office_file(tmp_path))
    issues = []
    items, meta = docx_items(source, issues)
    repeats = [i for i in items if i.text == "Repeat"]
    assert len(repeats) == 2 and repeats[0].locator != repeats[1].locator
    assert all(i.locator["kind"] == "docx_block" for i in items)
    assert {i.text for i in items} >= {
        "Header exact 987",
        "Footer exact 654",
        "Exact section",
    }
    assert {Path(p).name for p in meta["parts"]} == {
        "document.xml",
        "header1.xml",
        "footer1.xml",
    }
    pictures = [i for i in items if i.kind == "picture"]
    assert len(pictures) == 2 and all(
        i.asset_bytes == pictures[0].asset_bytes for i in pictures
    )
    assert any("anchor[1]" in i.locator["path"] for i in pictures)
    table = next(i for i in items if i.kind == "table")
    assert table.cells[0][0]["col_span"] == 2
    assert table.cells[1][0]["row_span"] == 2 and table.cells[2][0]["value"] is None
    snapshot = map_result(result_for(SourceFormat.DOCX, items, issues), source)
    assert all(
        n.annotation.provenance.producer_id == "producer-native-ooxml"
        for n in snapshot.structure[1:]
    )
    for cell in next(n for n in snapshot.structure if n.kind == "table").rows[1]:
        assert cell.annotation.provenance.evidence_ids
    output, *_ = output_bundle(snapshot)
    assert output.version == "1.2.0"


def workbook_file(root):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data TR"
    sheet.append(["Category", "Rooms", "Formula"])
    sheet.append(["Accessible", 7, "=B2*2"])
    sheet.append(["Total", 724, None])
    sheet["B2"].number_format = '0 "rooms"'
    sheet["A5"] = datetime(2026, 10, 2, tzinfo=UTC).replace(tzinfo=None)
    sheet["A5"].number_format = "yyyy-mm-dd"
    sheet.merge_cells("A7:B7")
    sheet["A7"] = "Merged"
    chart = BarChart()
    chart.title = "Native rooms"
    chart.y_axis.title = "Count"
    chart.add_data(
        Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True
    )
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
    sheet.add_chart(chart, "E2")
    path = root / "chart.xlsx"
    workbook.save(path)
    return path


def test_xlsx_typed_number_format_formula_missing_cache_and_native_chart(tmp_path):
    source = verified(workbook_file(tmp_path))
    issues = []
    wb = load_workbook(source.path)
    cache = load_workbook(source.path, data_only=True)
    charts = xlsx_charts(source, wb, cache, issues)
    assert len(charts) == 1
    chart = charts[0]
    assert chart.source_data["title"] == "Native rooms"
    assert chart.source_data["types"] == ["barChart"]
    ser = chart.source_data["series"][0]
    assert [c["value"] for c in ser["val"]["cells"]] == [7, 724]
    assert [c["value"] for c in ser["cat"]["cells"]] == ["Accessible", "Total"]
    assert ser["val"]["locator"] == {
        "kind": "spreadsheet_range",
        "sheet": "Data TR",
        "a1_range": "B2:B3",
    }
    assert ser["val"]["cells"][0]["number_format"] == '0 "rooms"'
    cell = _xlsx_cell(
        wb.active, cache.active, wb.active["C2"], issues, native_fidelity=True
    )
    assert (
        cell["formula"] == "=B2*2"
        and cell["cached_value"] is None
        and cell["value"] is None
    )
    assert any(i.code == "missing_cached_value" for i in issues)
    date = _xlsx_cell(
        wb.active, cache.active, wb.active["A5"], issues, native_fidelity=True
    )
    assert (
        date["value"] == "2026-10-02T00:00:00"
        and date["source_attributes"]["data_type"] == "d"
    )
    merged = _xlsx_cell(
        wb.active, cache.active, wb.active["B7"], issues, native_fidelity=True
    )
    assert merged["source_attributes"]["merge_covered"] is True
    snapshot = map_result(result_for(SourceFormat.XLSX, charts, issues), source)
    node = next(n for n in snapshot.structure if n.kind == "chart")
    assert node.field_annotations["/source_data/series/0/val"].provenance.evidence_ids
    output, _, _, files = output_bundle(snapshot)
    assert output.content[1].source_data["series"][0]["val"]["cells"][0]["value"] == 7
    jsonschema.Draft202012Validator(json.loads(files["ai.schema.json"])).validate(
        json.loads(files["ai.json"])
    )
    assert (
        output.quality.text_only_complete is False
    )  # native data is not semantic acceptance
    wb.close()
    cache.close()


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
@pytest.mark.parametrize("crop", [False, True])
def test_pdf_generic_misassigned_numeric_column_corrected_with_evidence(
    tmp_path, rotation, crop
):
    path = tmp_path / "grid.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page(width=600, height=800)
    for x in (80, 250, 350, 450):
        page.draw_line((x, 150), (x, 250))
    for y in (150, 200, 250):
        page.draw_line((80, y), (450, y))
    page.insert_text((90, 175), "ROOM")
    page.insert_text((260, 175), "LAND")
    page.insert_text((360, 175), "ACCESSIBLE")
    page.insert_text((90, 225), "Total")
    page.insert_text((360, 225), "7")
    if crop:
        page.set_cropbox(pymupdf.Rect(20, 20, 580, 760))
    page.set_rotation(rotation)
    pdf.save(path)
    pdf.close()
    source = verified(path)
    with pymupdf.open(path) as pdf:
        page = pdf[0]
        x0, y0, x1, y1 = page.find_tables(strategy="lines_strict").tables[0].bbox
        page.set_rotation(0)
        x0, y0, x1, y1 = page.find_tables(strategy="lines_strict").tables[0].bbox
        size = (page.cropbox.width, page.cropbox.height)
    item = StructuralItem(
        "table",
        "synthetic-grid",
        {
            "kind": "pdf_raw",
            "page_number": 1,
            "bbox": {"l": x0, "t": y0, "r": x1, "b": y1, "coord_origin": "TOPLEFT"},
        },
        page_size=size,
        cells=[
            [{"value": v} for v in ("ROOM", "LAND", "ACCESSIBLE")],
            [{"value": v} for v in ("Total", "7", "")],
        ],
    )
    items = [item]
    issues = []
    meta = reconcile_pdf_tables(source, items, issues, {})
    assert item.cells[1][1]["value"] == "" and item.cells[1][2]["value"] == "7"
    assert item.cells[1][1]["source_attributes"]["parser_text"] == "7"
    assert len(meta["changes"]) == 2 and all(
        v["anchor"] == "synthetic-grid" for v in meta["changes"]
    )
    snapshot = map_result(result_for(SourceFormat.PDF, items, issues), source)
    table = next(n for n in snapshot.structure if n.kind == "table")
    evidence = next(
        e
        for e in snapshot.evidence
        if e.id == table.rows[1][2].annotation.provenance.evidence_ids[0]
    )
    assert evidence.locator.bbox is not None
    assert (
        snapshot.knowledge_revision.coverage == "partial"
        and table.rows[1][2].annotation.review_status == "unreviewed"
    )


def test_old_schemas_stay_exact_and_new_facts_versioned():
    root = Path("packages/domain/docgrain_domain/canonical/schemas")
    for version in ("0.2.0", "0.3.0", "0.4.0", "0.5.0"):
        assert generated_core_schema_text(version) == (
            root / f"canonical-knowledge-{version}.schema.json"
        ).read_text(encoding="utf-8")
    for version in ("1.0.0", "1.1.0"):
        assert json.dumps(output_schema(version), ensure_ascii=False, sort_keys=True, indent=2) == Path(
            f"tests/fixtures/canonical/ai-document-{version}.schema.json"
        ).read_text(encoding="utf-8")


@pytest.mark.parametrize(
    "reference",
    [
        "'[remote.xlsx]Data TR'!$B$2:$B$3",
        "'Data TR'!$B$2:$B$20000",
        "'Missing'!$B$2:$B$3",
    ],
)
def test_external_or_oversized_chart_reference_stays_unresolved(tmp_path, reference):
    path = workbook_file(tmp_path)
    with ZipFile(path) as package:
        parts = {n: package.read(n) for n in package.namelist()}
    parts["xl/charts/chart1.xml"] = parts["xl/charts/chart1.xml"].replace(
        b"'Data TR'!$B$2:$B$3", reference.encode()
    )
    with ZipFile(path, "w") as package:
        for name, data in parts.items():
            package.writestr(name, data)
    wb = load_workbook(path)
    cached = load_workbook(path, data_only=True)
    issues = []
    chart = xlsx_charts(verified(path), wb, cached, issues)[0]
    assert chart.source_data["series"][0]["val"]["formula"] == reference
    assert chart.source_data["series"][0]["val"]["cells"] == []
    assert any(i.code == "chart_reference_unresolved" for i in issues)
    assert "/source_data/series/0/val" not in chart.field_locators
    wb.close()
    cached.close()


def test_docx_missing_header_and_field_result_are_explicit(tmp_path):
    from docgrain_worker.native_office import NS, R, W

    path = office_file(tmp_path)
    with ZipFile(path) as package:
        parts = {n: package.read(n) for n in package.namelist()}
    root = ET.fromstring(parts["word/document.xml"])
    root.find(".//w:headerReference", NS).set(R + "id", "missing-ref")
    p = root.find("w:body/w:p", NS)
    field = ET.SubElement(p, W + "fldSimple", {W + "instr": " PAGE "})
    ET.SubElement(ET.SubElement(field, W + "r"), W + "t").text = "99"
    parts["word/document.xml"] = ET.tostring(root)
    with ZipFile(path, "w") as package:
        for name, data in parts.items():
            package.writestr(name, data)
    issues = []
    items, meta = docx_items(verified(path), issues)
    assert not any(i.text == "Header exact 987" for i in items)
    assert {"part_reference_unresolved", "field_result_unverified"} <= {
        i.code for i in issues
    }
    assert meta["fields"][0]["visible_result"] == "99"
