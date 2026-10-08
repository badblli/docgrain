from datetime import date
from hashlib import sha256

import pytest
from docgrain_worker.structural import VerifiedSource, _xlsx_items
from docgrain_worker.xlsx_format import display_value
from openpyxl import Workbook


@pytest.mark.parametrize("value,fmt,expected", [
    (0.15, "0%", "15%"), (0.156, "0.0%", "15.6%"),
    (1234.565, "#,##0.00", "1,234.57"), (-42.5, "0.00", "-42.50"),
    (date(2026, 1, 2), "yyyy-mm-dd", "2026-01-02"),
    (date(2026, 1, 2), "dd.mm.yyyy", "02.01.2026"),
])
def test_number_percent_and_date_display(value, fmt, expected):
    assert display_value(value, fmt) == expected


def test_only_docling_cells_receive_display_and_formula_metadata(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Values"
    sheet.append([0.15, date(2026, 1, 2), "=1+2"])
    sheet["A1"].number_format = "0%"
    sheet["B1"].number_format = "yyyy-mm-dd"
    sheet["Z99"] = "Omitted sentinel"
    workbook.create_sheet("Omitted sheet")["A1"] = "Another sentinel"
    path = tmp_path / "sample.xlsx"
    workbook.save(path)
    workbook.close()
    data = path.read_bytes()
    raw = {"groups": [{"self_ref": "#/groups/0", "label": "sheet", "name": "Values"}],
           "tables": [{"self_ref": "#/tables/0", "parent": {"$ref": "#/groups/0"},
                       "prov": [{"bbox": {"l": 0, "t": 0, "r": 3, "b": 1, "coord_origin": "TOPLEFT"}}],
                       "data": {"num_rows": 1, "num_cols": 3, "table_cells": [
                           {"start_row_offset_idx": 0, "start_col_offset_idx": c, "text": text}
                           for c, text in enumerate(["0.15", "2026-01-02 00:00:00", ""]) ]}}]}
    issues = []
    items = _xlsx_items(VerifiedSource(path, sha256(data).hexdigest(), len(data)), raw, issues)
    tables = [item for item in items if item.kind == "table"]
    assert len(tables) == 1 and tables[0].locator["a1_range"] == "A1:C1"
    cells = tables[0].cells[0]
    assert cells[0]["value"] == 0.15 and cells[0]["display_text"] == "15%"
    assert cells[1]["display_text"] == "2026-01-02"
    assert cells[2]["formula"] == "=1+2" and cells[2]["cached_value"] is None
    assert cells[2]["value"] is None and cells[2]["display_text"] == ""
    assert {i.code for i in issues} == {"sheet_missing_in_docling", "missing_cached_value"}
    assert all(item.locator["sheet"] == "Values" for item in items)


def test_missing_docling_cell_inside_range_is_not_completed(tmp_path):
    workbook = Workbook()
    workbook.active.append(["Read", "Omitted"])
    path = tmp_path / "sample.xlsx"
    workbook.save(path)
    workbook.close()
    data = path.read_bytes()
    raw = {"groups": [{"self_ref": "#/groups/0", "label": "sheet", "name": "Sheet"}],
           "tables": [{"parent": {"$ref": "#/groups/0"},
                       "prov": [{"bbox": {"l": 0, "t": 0, "r": 2, "b": 1, "coord_origin": "TOPLEFT"}}],
                       "data": {"num_rows": 1, "num_cols": 2, "table_cells": [
                           {"start_row_offset_idx": 0, "start_col_offset_idx": 0, "text": "Read"}]}}]}
    items = _xlsx_items(VerifiedSource(path, sha256(data).hexdigest(), len(data)), raw, [])
    assert next(item for item in items if item.kind == "table").cells[0][1]["value"] == ""
