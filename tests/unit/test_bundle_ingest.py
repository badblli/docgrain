"""Exercise the CLI through a fake HTTP transport and the actual API routes."""

import json
from hashlib import sha256

import httpx
import pytest
from docgrain_api.main import app
from docgrain_api.routers import documents
from docgrain_domain import JobStatus, VersionStatus
from docgrain_domain.storage_paths import source_version_id, upload_key
from docgrain_ingest.cli import ingest_folder, main
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image


@pytest.fixture
def bundle_api(monkeypatch, live_repository):
    stored, uploads, queued = {}, [], []
    state = {"status": "done", "requests": []}

    def put(key, stream, mime, length):
        stored[key] = stream.read()
        uploads.append(key)
        assert len(stored[key]) == length

    def enqueue(job_id):
        queued.append(job_id)
        job = live_repository["jobs"][job_id]
        job.status = JobStatus(state["status"])
        version = live_repository["versions"][job.document_version_id]
        version.page_count = 2
        if state["status"] in {"done", "partial", "failed"}:
            version.status = VersionStatus(state["status"])
        if state["status"] == "partial":
            job.stages[2].attributes = {"structural_issues": [{"reason": "synthetic issue"}]}
            job.stages[3].attributes = job.stages[2].attributes
        if state["status"] == "failed":
            job.stages[2].error = "synthetic failure"

    monkeypatch.setattr(documents, "put_upload", put)
    monkeypatch.setattr(documents, "object_exists", lambda key: key in stored)
    monkeypatch.setattr(documents, "enqueue", enqueue)
    api = TestClient(app)

    def transport(request):
        state["requests"].append((request.method, request.url.path))
        response = api.request(request.method, str(request.url), content=request.content,
                               headers=dict(request.headers))
        return httpx.Response(response.status_code, content=response.content,
                              headers=response.headers, request=request)

    with httpx.Client(base_url="http://fake-api", transport=httpx.MockTransport(transport)) as client:
        yield client, api, live_repository, stored, uploads, queued, state


def synthetic_folder(root):
    root.mkdir()
    nested = root / "nested"
    nested.mkdir()
    (root / "guide.txt").write_text("Breakfast starts at 07:00.", encoding="utf-8")
    (nested / "copy.TXT").write_bytes((root / "guide.txt").read_bytes())
    (root / "sample.pdf").write_bytes(b"%PDF-synthetic")
    document = Document()
    document.add_paragraph("A neutral room description.")
    document.save(root / "rooms.docx")
    workbook = Workbook()
    workbook.active.append(["Room", "Capacity"])
    workbook.active.append(["Standard", 2])
    workbook.save(root / "rooms.xlsx")
    for name, format_name in [("map.png", "PNG"), ("photo.jpg", "JPEG"), ("photo.jpeg", "JPEG")]:
        Image.new("RGB", (3, 3), "white").save(root / name, format=format_name)
    (root / "ignored.bin").write_bytes(b"unsupported")
    return root


def test_rerun_reuses_content_and_workspaces_remain_separate(tmp_path, bundle_api):
    client, api, records, stored, uploads, queued, _ = bundle_api
    folder = synthetic_folder(tmp_path / "company")
    report_path = tmp_path / "bundle.json"
    first = ingest_folder(folder, "ws_alpha", client, report_path=report_path, poll_interval=0)
    assert len(first["files"]) == 9
    assert sum(row["status"] == "done" for row in first["files"]) == 8
    assert next(row for row in first["files"] if row["file"] == "ignored.bin")["issues"]
    # TXT copy and JPEG alias are the same content, irrespective of filename.
    assert len(records["documents"]) == 6
    assert len(uploads) == len(queued) == 6
    ids = [row["document_id"] for row in first["files"]]
    second = ingest_folder(folder, "ws_alpha", client, report_path=report_path, poll_interval=0)
    assert [row["document_id"] for row in second["files"]] == ids
    assert all(row["reused"] for row in second["files"] if row["status"] == "done")
    assert len(records["documents"]) == len(uploads) == len(queued) == 6
    assert json.loads(report_path.read_text(encoding="utf-8")) == second
    other = ingest_folder(folder, "ws_beta", client, report_path=report_path, poll_interval=0)
    assert {row["document_id"] for row in other["files"] if row["document_id"]}.isdisjoint(
        row["document_id"] for row in first["files"] if row["document_id"])
    assert len(records["documents"]) == 12
    listed = api.get("/v1/documents", params={"workspace_id": "ws_alpha", "limit": 100}).json()
    assert len(listed) == 6
    assert all(item["document"]["workspace_id"] == "ws_alpha" for item in listed)
    assert all(key.startswith(("uploads/ws_alpha/", "uploads/ws_beta/")) for key in stored)
    assert all(row["pages"] == 2 for row in first["files"] if row["status"] == "done")


@pytest.mark.parametrize("uploaded", [False, True])
def test_interrupted_upload_reuses_registration_and_original_filename(tmp_path, bundle_api, uploaded):
    client, api, records, _, uploads, queued, _ = bundle_api
    folder = tmp_path / "company"
    folder.mkdir()
    data = b"Neutral source text."
    (folder / "renamed.txt").write_bytes(data)
    original = api.post("/v1/documents", json={
        "workspace_id": "ws_alpha", "filename": "original.txt", "mime_type": "text/plain",
        "byte_size": len(data), "content_sha256": sha256(data).hexdigest(),
    }).json()
    if uploaded:
        api.put(original["upload_url"], files={"file": ("original.txt", data, "text/plain")})
    report = ingest_folder(folder, "ws_alpha", client, report_path=tmp_path / "report.json", poll_interval=0)
    row = report["files"][0]
    assert row["reused"] and row["status"] == "done"
    assert row["document_id"] == original["document"]["id"]
    assert len(records["documents"]) == len(uploads) == len(queued) == 1


@pytest.mark.parametrize("status", ["partial", "failed", "queued"])
def test_processing_issues_and_timeout_are_honest(tmp_path, bundle_api, status):
    client, _, records, _, _, _, state = bundle_api
    state["status"] = status
    folder = tmp_path / "company"
    folder.mkdir()
    (folder / "guide.txt").write_text("Neutral text.", encoding="utf-8")
    report = ingest_folder(folder, "ws_alpha", client, report_path=tmp_path / "report.json",
                           timeout=0.001, poll_interval=0)
    row = report["files"][0]
    assert row["status"] == ("timeout" if status == "queued" else status)
    assert len(row["issues"]) == 1
    assert row["document_id"] in records["documents"]


def test_invalid_file_does_not_stop_rest_of_folder(tmp_path, bundle_api):
    client, _, _, _, _, _, _ = bundle_api
    folder = tmp_path / "company"
    folder.mkdir()
    (folder / "bad.pdf").write_bytes(b"not a PDF")
    (folder / "empty.txt").touch()
    (folder / "good.txt").write_text("Neutral text.", encoding="utf-8")
    report = ingest_folder(folder, "ws_alpha", client, report_path=tmp_path / "report.json", poll_interval=0)
    assert [row["status"] for row in report["files"]] == ["error", "error", "done"]
    assert report["files"][0]["issues"] == [{"reason": "HTTP 415"}]


def test_transport_failure_is_reported_per_file(tmp_path):
    folder = tmp_path / "company"
    folder.mkdir()
    (folder / "guide.txt").write_text("Neutral text.", encoding="utf-8")

    def fail(request):
        raise httpx.ConnectError("synthetic unavailable endpoint", request=request)

    with httpx.Client(base_url="http://fake", transport=httpx.MockTransport(fail)) as client:
        report = ingest_folder(folder, "ws_alpha", client, report_path=tmp_path / "report.json")
    assert report["files"][0]["status"] == "error"
    assert report["files"][0]["issues"] == [{"reason": "ConnectError"}]


def test_cli_exit_code_and_report_excluded_on_rerun(tmp_path, monkeypatch, bundle_api):
    client, _, _, _, _, _, state = bundle_api
    client_type = type(client)
    monkeypatch.setattr("docgrain_ingest.cli.httpx.Client", lambda **kwargs:
                        client_type(base_url="http://fake-api", transport=client._transport))
    folder = tmp_path / "company"
    folder.mkdir()
    (folder / "guide.txt").write_text("Neutral text.", encoding="utf-8")
    args = ["ingest-folder", str(folder), "--workspace", "ws_alpha", "--report", str(folder / "report.json")]
    assert main(args) == 0
    report = json.loads((folder / "report.json").read_text(encoding="utf-8"))
    assert len(report["files"]) == 1
    state["status"] = "partial"
    (folder / "new.txt").write_text("Different neutral text.", encoding="utf-8")
    assert main(args) == 1
    assert len(json.loads((folder / "report.json").read_text(encoding="utf-8"))["files"]) == 2


def test_workspace_storage_paths_accept_legacy_and_reject_other_workspace():
    assert upload_key("ws_alpha", "doc", "ver") == "uploads/ws_alpha/doc/ver/original"
    assert source_version_id("s3://bucket/uploads/doc/ver/original", "ws_alpha", "doc") == "ver"
    assert source_version_id("s3://bucket/uploads/ws_alpha/doc/ver/original", "ws_alpha", "doc") == "ver"
    assert source_version_id("s3://bucket/uploads/ws_beta/doc/ver/original", "ws_alpha", "doc") is None


def test_registration_rejects_invalid_workspace_and_checksum(bundle_api):
    _, api, records, _, _, _, _ = bundle_api
    payload = {"workspace_id": "../unsafe", "filename": "guide.txt", "mime_type": "text/plain", "byte_size": 1}
    assert api.post("/v1/documents", json=payload).status_code == 422
    payload.update(workspace_id="ws_alpha", content_sha256="z" * 64)
    assert api.post("/v1/documents", json=payload).status_code == 422
    assert not records["documents"]
